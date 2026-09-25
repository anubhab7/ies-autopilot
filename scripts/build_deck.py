# Builds docs/IES_Autopilot_Deck.pptx from docs/DECK_OUTLINE.md content, then verifies it.
# Usage: python3 scripts/build_deck.py  (requires python-pptx)
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, "docs", "screenshots")
OUT = os.path.join(ROOT, "docs", "IES_Autopilot_Deck.pptx")

BG = RGBColor(0x0D, 0x15, 0x24)
TEXT = RGBColor(0xE6, 0xED, 0xF7)
MUTED = RGBColor(0x8F, 0xA3, 0xBF)
ACCENT = RGBColor(0x4C, 0x8D, 0xFF)
HAIRLINE = RGBColor(0x25, 0x35, 0x4F)
RAISED = RGBColor(0x1A, 0x27, 0x40)
FONT = "Calibri"

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.4
GAP = 0.4
HALF_W = (SLIDE_W - 2 * MARGIN - GAP) / 2
RIGHT_X = MARGIN + HALF_W + GAP
FOOTER_TEXT = "IES Autopilot | Anubhab Chakraborty | Intuit PM Intern case"
SHOT_RATIO = 900 / 1440

PROTOTYPE_URL = "https://ies-autopilot.b26015.workers.dev"
GITHUB_URL = "https://github.com/anubhab7/ies-autopilot"

# Each body paragraph is a list of (text, bold) runs.
SLIDES = [
    {
        "title": "The customer problem",
        "body": [
            [("Who: ", True), ("Maya, a Controller at a 420-person, three-entity company, is blocked by intercompany mismatches, manual reconciliations, and agents she cannot trust.", False)],
            [("Goal: ", True), ("close in 3 days. Only 18% of teams do; half take 6 or more (Ledge).", False)],
        ],
        "shots": ["02-morning-brief.png"],
    },
    {
        "title": "What IES has today, and the gap",
        "body": [
            [("Today: ", True), ("7 function agents, used monthly by 75% of customers.", False)],
            [("Gaps: ", True), ("agents map to functions, not outcomes; trust tools are missing; developers pay for data and wait 30 days for review.", False)],
        ],
        "shots": ["17-research.png"],
    },
    {
        "title": "Vision and three pillars",
        "body": [
            [("Every mid-market finance team gets an AI crew with an expert on call.", False)],
            [("1. ", True), ("Outcome Autopilots, not feature agents.", False)],
            [("2. ", True), ("Trust by design, with guardrails and undo.", False)],
            [("3. ", True), ("An ecosystem where builders keep 80% to 85%.", False)],
        ],
        "shots": ["16-strategy.png"],
    },
    {
        "title": "Customer experience",
        "body": [
            [("Close Autopilot: ", True), ("one board, four lanes, a 9-day close cut to 3.", False)],
            [("Exceptions: ", True), ("evidence, a balanced draft, and why each landed in its lane.", False)],
            [("Expert on call: ", True), ("Priya Raman replies with a reviewed entry.", False)],
        ],
        "shots": ["03-close-autopilot.png", "04-exception-ic-310.png", "07-expert-reply.png"],
    },
    {
        "title": "Human plus AI operating model",
        "body": [
            [("Rule: ", True), ("agents post alone only when confident, reversible, and strictly below materiality.", False)],
            [("Example: ", True), ("at a $40,000 limit, FX-77 ($38,100) auto-posts; DEP-5 ($40,000) waits for Maya.", False)],
            [("Autonomy Dial: ", True), ("L0 to L3 per workflow, previewed first.", False)],
        ],
        "shots": ["05-control-tower.png"],
    },
    {
        "title": "Developer journey",
        "body": [
            [("Onboard: ", True), ("Hangar sandbox with synthetic data and a first-call timer.", False)],
            [("Build: ", True), ("SDK, no-code, REST, and hosted MCP that drafts, never posts.", False)],
            [("Certify and earn: ", True), ("50 evals, v2 scores 98%; live in 3 days, keep 80%.", False)],
        ],
        "shots": ["12-hangar.png", "14-studio-evals.png", "15-publish.png"],
    },
    {
        "title": "Business and ecosystem model",
        "body": [
            [("Revenue: ", True), ("outcome-priced Autopilot tier, expert session fees, and an Agent Store take rate.", False)],
            [("Illustrative GMV: ", True), ("8,000 x 25% x 2 x $250 x 12 = $12M; $2.4M to Intuit at 20%.", False)],
            [("Free sandbox and reads attract builders.", False)],
        ],
        "shots": ["10-agent-detail-ledgerloop.png"],
    },
    {
        "title": "Roadmap, experiments, and KPIs",
        "body": [
            [("Now: ", True), ("Close Autopilot and the dial. ", False), ("Next: ", True), ("Agent Store. ", False), ("Later: ", True), ("outcome pricing.", False)],
            [("Why: ", True), ("prove trust on the close before third-party agents act.", False)],
            [("North Star: ", True), ("verified agent hours, tested first with a Wizard-of-Oz pilot.", False)],
        ],
        "shots": ["11-flight-log.png"],
    },
]

PLACEHOLDER = "Edit me"
AI_TABLE = [
    ("Stage", "AI tool", "Example prompt"),
    ("Empathize", PLACEHOLDER, PLACEHOLDER),
    ("Define", PLACEHOLDER, PLACEHOLDER),
    ("Ideate", PLACEHOLDER, PLACEHOLDER),
    ("Prototype", "Claude Code", "Build IES Autopilot in 9 phases from BUILD_PROMPT.md, with seed data and tests."),
    ("Experiment", PLACEHOLDER, PLACEHOLDER),
]


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def style_run(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def add_title(slide, text, width=HALF_W):
    box = slide.shapes.add_textbox(Inches(MARGIN), Inches(MARGIN), Inches(width), Inches(1.2))
    box.name = "Title"
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    style_run(run, 32, TEXT, bold=True)
    return box


def add_body(slide, paragraphs, top=1.8, width=HALF_W, height=4.7):
    box = slide.shapes.add_textbox(Inches(MARGIN), Inches(top), Inches(width), Inches(height))
    box.name = "Body"
    tf = box.text_frame
    tf.word_wrap = True
    for i, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.JUSTIFY
        p.space_after = Pt(14)
        p.line_spacing = 1.15
        for text, bold in runs:
            run = p.add_run()
            run.text = text
            style_run(run, 16, ACCENT if bold else TEXT, bold=bold)
    return box


def add_footer(slide):
    box = slide.shapes.add_textbox(Inches(MARGIN), Inches(SLIDE_H - MARGIN - 0.3), Inches(SLIDE_W - 2 * MARGIN), Inches(0.3))
    box.name = "Footer"
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = FOOTER_TEXT
    style_run(run, 9, MUTED)


def add_picture(slide, name, x, y, w):
    path = os.path.join(SHOTS, name)
    if not os.path.exists(path):
        raise FileNotFoundError(f"Missing screenshot: {path}")
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    pic.line.color.rgb = HAIRLINE
    pic.line.width = Pt(0.75)
    return pic


def add_shots(slide, shots):
    content_top = MARGIN
    content_h = SLIDE_H - 2 * MARGIN - 0.45
    if len(shots) == 1:
        h = HALF_W * SHOT_RATIO
        add_picture(slide, shots[0], RIGHT_X, content_top + (content_h - h) / 2, HALF_W)
        return
    # One large screenshot on top, two smaller ones side by side below it.
    inner_gap = 0.15
    big_h = HALF_W * SHOT_RATIO
    small_w = (HALF_W - inner_gap) / 2
    small_h = small_w * SHOT_RATIO
    total = big_h + inner_gap + small_h
    y = content_top + (content_h - total) / 2
    add_picture(slide, shots[0], RIGHT_X, y, HALF_W)
    add_picture(slide, shots[1], RIGHT_X, y + big_h + inner_gap, small_w)
    add_picture(slide, shots[2], RIGHT_X + small_w + inner_gap, y + big_h + inner_gap, small_w)


def build_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)

    box = slide.shapes.add_textbox(Inches(MARGIN), Inches(1.6), Inches(HALF_W), Inches(1.2))
    box.name = "Title"
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "IES Autopilot"
    style_run(run, 54, ACCENT, bold=True)

    body = slide.shapes.add_textbox(Inches(MARGIN), Inches(2.9), Inches(HALF_W), Inches(3.6))
    body.name = "Body"
    tf = body.text_frame
    tf.word_wrap = True

    def para(first=False):
        para_obj = tf.paragraphs[0] if first else tf.add_paragraph()
        para_obj.alignment = PP_ALIGN.JUSTIFY
        para_obj.space_after = Pt(10)
        return para_obj

    p = para(first=True)
    p.space_after = Pt(28)
    r = p.add_run()
    r.text = "Your finance team's AI crew, with a human expert always on call"
    style_run(r, 20, TEXT)

    p = para()
    p.space_after = Pt(2)
    r = p.add_run()
    r.text = "Anubhab Chakraborty"
    style_run(r, 16, TEXT, bold=True)

    p = para()
    p.space_after = Pt(24)
    r = p.add_run()
    r.text = "Aspiring product manager who designed and built this prototype for the Intuit PM Intern case."
    style_run(r, 16, MUTED)

    for label, url in (("Live prototype: ", PROTOTYPE_URL), ("GitHub: ", GITHUB_URL)):
        p = para()
        r = p.add_run()
        r.text = label
        style_run(r, 16, MUTED)
        r = p.add_run()
        r.text = url
        style_run(r, 16, ACCENT)
        r.hyperlink.address = url
        r.font.underline = True

    add_picture(slide, "01-home.png", RIGHT_X, (SLIDE_H - HALF_W * SHOT_RATIO) / 2, HALF_W)


def build_content(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, spec["title"])
    add_body(slide, spec["body"])
    add_shots(slide, spec["shots"])
    add_footer(slide)


def build_ai_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    table_w = 7.9
    add_title(slide, "How I used AI", width=table_w)

    rows, cols = len(AI_TABLE), 3
    shape = slide.shapes.add_table(rows, cols, Inches(MARGIN), Inches(1.5), Inches(table_w), Inches(3.6))
    shape.name = "Body table"
    table = shape.table
    for idx, w in enumerate((1.5, 1.6, table_w - 3.1)):
        table.columns[idx].width = Inches(w)
    for r, row in enumerate(AI_TABLE):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RAISED if r == 0 else BG
            cell.margin_left = cell.margin_right = Inches(0.1)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.JUSTIFY
            run = p.add_run()
            run.text = value
            header = r == 0
            color = MUTED if value == PLACEHOLDER else (ACCENT if header else TEXT)
            style_run(run, 14, color, bold=header or c == 0)

    add_body(
        slide,
        [[("Next steps: ", True), ("run four rapid tests, interview 8 controllers, and pilot 3 partner agents.", False)]],
        top=5.4,
        width=table_w,
        height=1.0,
    )
    shot_w = SLIDE_W - MARGIN - (MARGIN + table_w + GAP)
    add_picture(slide, "18-process.png", MARGIN + table_w + GAP, 1.5, shot_w)
    add_footer(slide)


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    build_cover(prs)
    for spec in SLIDES:
        build_content(prs, spec)
    build_ai_slide(prs)
    prs.save(OUT)


def iter_paragraphs(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                yield shape.name, p
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        yield shape.name, p


def verify():
    prs = Presentation(OUT)
    dashes = (chr(0x2013), chr(0x2014))
    results = {}

    bad_runs = [
        (i + 1, r.text)
        for i, s in enumerate(prs.slides)
        for _, p in iter_paragraphs(s)
        for r in p.runs
        if any(d in r.text for d in dashes)
    ]
    results["No em or en dashes in any text run"] = not bad_runs

    bad_align = [
        (i + 1, name, "".join(r.text for r in p.runs))
        for i, s in enumerate(prs.slides)
        for name, p in iter_paragraphs(s)
        if name.startswith("Body") and p.alignment != PP_ALIGN.JUSTIFY
    ]
    results["Every body paragraph is justified"] = not bad_align

    bad_fill = [
        i + 1
        for i, s in enumerate(prs.slides)
        if s.background.fill.type != MSO_FILL.SOLID or s.background.fill.fore_color.rgb != BG
    ]
    results["Every slide has a solid #0D1524 fill"] = not bad_fill

    results["Exactly 10 slides"] = len(prs.slides) == 10

    for label, ok in results.items():
        print(f"{'PASS' if ok else 'FAIL'}: {label}")
    for detail in (bad_runs, bad_align, bad_fill):
        if detail:
            print("  details:", detail)

    for i, s in enumerate(prs.slides):
        words = sum(
            len("".join(r.text for r in p.runs).split())
            for name, p in iter_paragraphs(s)
            if name == "Body"
        )
        print(f"  slide {i + 1}: {words} body words")
    return all(results.values())


if __name__ == "__main__":
    build()
    sys.exit(0 if verify() else 1)
